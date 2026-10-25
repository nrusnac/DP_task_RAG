"""Format-specific parsers -> ordered structural elements.

Every parser returns (title, elements); an element is a dict:
  {"kind": "heading", "level": int, "text": str}
  {"kind": "para", "text": str}
  {"kind": "table", "rows": [[cell, ...], ...], "title": str | None}   # first row = header
  {"kind": "slide" | "email", "text": str, "title": str}

Edge cases found while exploring the corpus (see docs/01-dataset.md):
- docx tracked changes (w:ins / w:del) are rendered as [INSERTED: ...] / [DELETED: ...]; python-docx drops them.
- malformed docx XML (an unescaped '&') is parsed with lxml's recover mode instead of crashing; malformed xlsx
  XML has its bare '&' escaped and is reloaded.
- xlsx formula cells without a cached value keep their formula text instead of becoming empty.
"""
from __future__ import annotations

import email
import re
import zipfile
from email import policy
from html.parser import HTMLParser
from pathlib import Path

from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_XML = etree.XMLParser(recover=True, huge_tree=True)
_HEADING_RE = re.compile(r"^(ARTICLE|Article|SECTION|Section|PART|Part|SCHEDULE|Schedule|EXHIBIT|Exhibit|ANNEX|Annex)\s+"
                         r"[\dIVXLC]+[A-Za-z]?\b")
_NUMBERED_RE = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+[A-Z][^.]{0,90}$")


def _ws(s: str) -> str:
    return re.sub(r"[ \t ]+", " ", s).strip()


# ----------------------------------------------------------------------------------------------- docx
def _docx_styles(z: zipfile.ZipFile) -> dict[str, tuple[str, int | None]]:
    """styleId -> (name, heading level or None)."""
    out = {}
    try:
        root = etree.fromstring(z.read("word/styles.xml"), _XML)
    except KeyError:
        return out
    for st in root.iter(W + "style"):
        sid = st.get(W + "styleId")
        name_el = st.find(W + "name")
        name = (name_el.get(W + "val") if name_el is not None else sid or "").lower()
        level = None
        m = re.match(r"heading\s*(\d)", name)
        if m:
            level = int(m.group(1))
        elif name in ("title",):
            level = 0
        out[sid] = (name, level)
    return out


def _run_state(run, para) -> str:
    node = run.getparent()
    while node is not None and node is not para:
        tag = etree.QName(node).localname
        if tag in ("ins", "moveTo"):
            return "ins"
        if tag in ("del", "moveFrom"):
            return "del"
        node = node.getparent()
    return "del" if run.find(W + "delText") is not None else "norm"


def _para_text(p) -> tuple[str, bool]:
    """Paragraph text with tracked changes made explicit. Returns (text, all_bold)."""
    pieces: list[tuple[str, str]] = []
    bold_flags = []
    for r in p.iter(W + "r"):
        buf = []
        for el in r:
            tag = etree.QName(el).localname
            if tag in ("t", "delText"):
                buf.append(el.text or "")
            elif tag == "tab":
                buf.append("\t")
            elif tag in ("br", "cr"):
                buf.append("\n")
        txt = "".join(buf)
        if not txt:
            continue
        rpr = r.find(W + "rPr")
        bold_flags.append(rpr is not None and rpr.find(W + "b") is not None and txt.strip() != "")
        state = _run_state(r, p)
        if pieces and pieces[-1][0] == state:
            pieces[-1] = (state, pieces[-1][1] + txt)
        else:
            pieces.append((state, txt))
    out = []
    for state, txt in pieces:
        if state == "ins":
            out.append(f"[INSERTED: {txt.strip()}]")
        elif state == "del":
            out.append(f"[DELETED: {txt.strip()}]")
        else:
            out.append(txt)
    return _ws("".join(out)), bool(bold_flags) and all(bold_flags)


def _para_level(p, styles) -> int | None:
    ppr = p.find(W + "pPr")
    if ppr is None:
        return None
    ps = ppr.find(W + "pStyle")
    if ps is not None and ps.get(W + "val") in styles:
        lvl = styles[ps.get(W + "val")][1]
        if lvl is not None:
            return lvl
    ol = ppr.find(W + "outlineLvl")
    if ol is not None and ol.get(W + "val", "").isdigit():
        return int(ol.get(W + "val")) + 1
    return None


def _table_rows(tbl) -> list[list[str]]:
    rows = []
    for tr in tbl.iter(W + "tr"):
        cells = []
        for tc in tr.findall(W + "tc"):
            cells.append(_ws(" ".join(_para_text(p)[0] for p in tc.iter(W + "p"))))
        if any(cells):
            rows.append(cells)
    return rows


def _iter_body(node):
    """Yield w:p and w:tbl in document order, descending into content controls (w:sdt)."""
    for child in node:
        tag = etree.QName(child).localname
        if tag in ("p", "tbl"):
            yield child
        elif tag in ("sdt", "sdtContent", "customXml"):
            yield from _iter_body(child)


def parse_docx(path: Path):
    z = zipfile.ZipFile(path)
    styles = _docx_styles(z)
    root = etree.fromstring(z.read("word/document.xml"), _XML)
    body = root.find(W + "body")
    elements, title = [], None
    for node in _iter_body(body if body is not None else root):
        if etree.QName(node).localname == "tbl":
            rows = _table_rows(node)
            if rows:
                elements.append({"kind": "table", "rows": rows, "title": None})
            continue
        text, all_bold = _para_text(node)
        if not text:
            continue
        level = _para_level(node, styles)
        if level is None and len(text) < 110 and (_HEADING_RE.match(text) or (all_bold and not text.endswith((".", ":", ";")))
                                                 or (_NUMBERED_RE.match(text) and all_bold)):
            level = 2 if _HEADING_RE.match(text) else 3
        if level is not None:
            if level == 0 or (title is None and level <= 1):
                title = title or text
            elements.append({"kind": "heading", "level": max(level, 1), "text": text})
        else:
            elements.append({"kind": "para", "text": text})
    if title is None:
        first = next((e["text"] for e in elements if e["kind"] in ("heading", "para")), None)
        title = first[:120] if first else path.stem
    return title, elements


# ----------------------------------------------------------------------------------------------- xlsx
def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.4f}".rstrip("0").rstrip(".") if abs(v) < 1e15 else str(v)
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}" if abs(v) >= 10000 else str(v)
    return _ws(str(v))


_BARE_AMP = re.compile(rb"&(?!(?:[A-Za-z]+|#\d+|#x[0-9A-Fa-f]+);)")


def _repaired_zip(path: Path):
    """Same workbook with bare '&' in its XML parts escaped (openpyxl has no recover mode)."""
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            zout.writestr(item, _BARE_AMP.sub(b"&amp;", data) if item.filename.endswith(".xml") else data)
    buf.seek(0)
    return buf


def parse_xlsx(path: Path):
    import openpyxl
    try:
        wv = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path, data_only=False)
    except Exception:  # noqa: BLE001  -- malformed XML (unescaped '&'): repair and retry once
        wv = openpyxl.load_workbook(_repaired_zip(path), data_only=True)
        wf = openpyxl.load_workbook(_repaired_zip(path), data_only=False)
    elements = []
    for ws in wv.worksheets:
        wsf = wf[ws.title]
        rows = []
        for r_v, r_f in zip(ws.iter_rows(), wsf.iter_rows()):
            cells = []
            for cv, cf in zip(r_v, r_f):
                val = cv.value
                if val is None and isinstance(cf.value, str) and cf.value.startswith("="):
                    val = f"[formula {cf.value}]"  # no cached value in the file
                cells.append(_fmt(val))
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        elements.append({"kind": "heading", "level": 1, "text": f"Sheet: {ws.title}"})
        # leading single-cell rows are sheet titles/notes, not the table header
        while len(rows) > 1 and sum(1 for c in rows[0] if c) == 1:
            elements.append({"kind": "para", "text": next(c for c in rows[0] if c)})
            rows = rows[1:]
        elements.append({"kind": "table", "rows": rows, "title": ws.title})
    return path.stem.replace("-", " ").replace("_", " ").title(), elements


# ----------------------------------------------------------------------------------------------- pptx
def parse_pptx(path: Path):
    from pptx import Presentation
    prs = Presentation(path)
    elements, deck_title = [], None
    for i, slide in enumerate(prs.slides, 1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = _ws(slide.shapes.title.text_frame.text) or None
        parts = []
        for sh in slide.shapes:
            if sh == slide.shapes.title:
                continue
            if sh.has_text_frame:
                t = "\n".join(_ws(p.text if hasattr(p, "text") else "") for p in sh.text_frame.paragraphs)
                if t.strip():
                    parts.append(t.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[_ws(c.text) for c in r.cells] for r in sh.table.rows]
                parts.append(rows_to_markdown(rows))
        if slide.has_notes_slide:
            notes = _ws(slide.notes_slide.notes_text_frame.text or "")
            if notes:
                parts.append(f"Speaker notes: {notes}")
        deck_title = deck_title or title
        stitle = f"Slide {i}" + (f": {title}" if title else "")
        elements.append({"kind": "slide", "title": stitle, "text": "\n".join(parts)})
    return deck_title or path.stem, elements


# ----------------------------------------------------------------------------------------------- eml/txt
class _Strip(HTMLParser):
    def __init__(self):
        super().__init__()
        self.out = []

    def handle_data(self, d):
        self.out.append(d)

    def handle_starttag(self, tag, attrs):
        if tag in ("br", "p", "div", "tr", "li"):
            self.out.append("\n")


def parse_eml(path: Path):
    m = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    body = m.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        s = _Strip(); s.feed(text); text = "".join(s.out)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    subject = str(m["subject"] or path.stem)
    hdr = f"From: {m['from']}\nTo: {m['to']}\n" + (f"Cc: {m['cc']}\n" if m["cc"] else "") + \
          f"Date: {m['date']}\nSubject: {subject}"
    return subject, [{"kind": "email", "title": f"Email: {subject}", "text": hdr + "\n\n" + text}]


def parse_txt(path: Path):
    text = path.read_text(errors="ignore")
    paras = [_ws(p) for p in re.split(r"\n\s*\n", text) if p.strip()]
    title = paras[0][:120] if paras else path.stem
    return title, [{"kind": "para", "text": p} for p in paras]


PARSERS = {".docx": parse_docx, ".xlsx": parse_xlsx, ".pptx": parse_pptx, ".eml": parse_eml, ".txt": parse_txt,
           ".md": parse_txt}


def parse(path: Path):
    fn = PARSERS.get(path.suffix.lower())
    if fn is None:
        raise ValueError(f"unsupported file type: {path.suffix}")
    return fn(path)


def rows_to_markdown(rows: list[list[str]], header: list[str] | None = None) -> str:
    if not rows and not header:
        return ""
    if header is None:
        header, body = rows[0], rows[1:]
    else:
        body = rows
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    pad = lambda r: [c.replace("|", "/") for c in r] + [""] * (width - len(r))  # noqa: E731
    lines = ["| " + " | ".join(pad(header)) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(pad(r)) + " |" for r in body]
    return "\n".join(lines)
